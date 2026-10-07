"""Import questions from the supplied 1C:Enterprise 8.3 workbook.

Column A starts each question, column B lists its options, and the bold option
is the answer key. Run with the workbook path as the first argument.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    source = Path(sys.argv[1])
    sheet = load_workbook(source, data_only=True).active
    groups = []
    current = None

    for row in range(2, sheet.max_row + 1):
        prompt = sheet.cell(row, 1).value
        option = sheet.cell(row, 2)
        if prompt is not None:
            current = {"question": str(prompt).strip(), "options": []}
            groups.append(current)
        if current is None or option.value is None:
            raise ValueError(f"Missing question or option at row {row}")
        current["options"].append((str(option.value).strip(), bool(option.font.bold)))

    questions = []
    for number, group in enumerate(groups, start=1):
        options = group["options"]
        answers = [index for index, (_, bold) in enumerate(options, start=1) if bold]
        if not group["question"] or len(options) < 2 or any(not text for text, _ in options) or len(answers) != 1:
            raise ValueError(f"Invalid question {number}")
        answer = answers[0]
        questions.append({
            "id": f"1.{number}",
            "section_number": 1,
            "section": "Вопросы по платформе",
            "question_number": number,
            "question": group["question"],
            "options": "\n".join(f"{index}) {text}" for index, (text, _) in enumerate(options, start=1)),
            "answer_number": answer,
            "answer_text": options[answer - 1][0],
            "confirmed": False,
        })

    target = Path(__file__).resolve().parents[1] / "1C_Enterprise83_questions.json"
    target.write_text(json.dumps(questions, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(questions)} questions into {target}")


if __name__ == "__main__":
    main()
